from __future__ import annotations

from io import BytesIO
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from django.conf import settings

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class CustomerStatementExportService:

    @staticmethod
    def _money(value) -> str:
        amount = Decimal(str(value or 0))
        return "$ " + f"{amount:,.0f}".replace(",", ".")

    @classmethod
    def build_pdf(cls, snapshot: dict) -> bytes:

        output = BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=10 * mm,
            bottomMargin=12 * mm,
            title="Estado de Cuenta",
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "StatementTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=17,
            textColor=colors.HexColor("#25312B"),
            spaceAfter=2,
        )

        customer_style = ParagraphStyle(
            "StatementCustomer",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=11,
            textColor=colors.HexColor("#34413B"),
        )

        meta_style = ParagraphStyle(
            "StatementMeta",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.8,
            leading=10,
            textColor=colors.HexColor("#707D76"),
        )

        normal_style = ParagraphStyle(
            "StatementCell",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.2,
            leading=8.5,
            alignment=TA_LEFT,
        )

        table_body_style = ParagraphStyle(
            "StatementTableBody",
            parent=normal_style,
            fontName="Helvetica",
            fontSize=7.2,
            leading=8.5,
            textColor=colors.HexColor("#222222"),
            alignment=TA_LEFT,
        )

        table_body_bold_style = ParagraphStyle(
            "StatementTableBodyBold",
            parent=table_body_style,
            fontName="Helvetica-Bold",
        )

        table_body_center_style = ParagraphStyle(
            "StatementTableBodyCenter",
            parent=table_body_style,
            alignment=TA_CENTER,
        )

        table_body_right_style = ParagraphStyle(
            "StatementTableBodyRight",
            parent=table_body_style,
            alignment=TA_RIGHT,
        )

        table_body_right_bold_style = ParagraphStyle(
            "StatementTableBodyRightBold",
            parent=table_body_right_style,
            fontName="Helvetica-Bold",
        )

        document_meta_style = ParagraphStyle(
            "StatementDocumentMeta",
            parent=normal_style,
            fontName="Helvetica",
            fontSize=7.2,
            leading=8.5,
            textColor=colors.HexColor("#222222"),
            spaceBefore=0,
        )

        refact_badge_style = ParagraphStyle(
            "StatementRefactBadge",
            parent=normal_style,
            fontName="Helvetica-Bold",
            fontSize=5.8,
            leading=6.4,
            textColor=colors.HexColor("#075D78"),
            alignment=TA_LEFT,
        )

        refact_reference_style = ParagraphStyle(
            "StatementRefactReference",
            parent=normal_style,
            fontName="Helvetica",
            fontSize=6.2,
            leading=7.3,
            textColor=colors.HexColor("#52616B"),
            spaceBefore=1,
        )

        claim_badge_style = ParagraphStyle(
            "StatementClaimBadge",
            parent=normal_style,
            fontName="Helvetica-Bold",
            fontSize=5.8,
            leading=6.4,
            textColor=colors.HexColor("#9A6700"),
            alignment=TA_LEFT,
        )

        section_style = ParagraphStyle(
            "StatementSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=11,
            textColor=colors.HexColor("#29362F"),
        )

        story = []

        customer = snapshot["customer"]
        totals = snapshot["totals"]

        # ====================================================
        # CABECERA
        # ====================================================

        header_copy = [
            Paragraph(
                "Estado de Cuenta",
                title_style,
            ),
            Paragraph(
                str(customer["name"]),
                customer_style,
            ),
            Paragraph(
                (
                    f"RUT {customer.get('rut') or '-'}"
                    f" · Fecha de emisión: "
                    f"{snapshot['as_of_date']}"
                ),
                meta_style,
            ),
        ]

        logo = ""

        try:
            from django.conf import settings
            from pathlib import Path
            from reportlab.platypus import Image

            logo_path = (
                Path(settings.BASE_DIR)
                / "static"
                / "img"
                / "brand"
                / "logo_mosaico.png"
            )

            if logo_path.exists():
                logo = Image(
                    str(logo_path),
                    width=43 * mm,
                    height=11 * mm,
                )
                logo.hAlign = "RIGHT"

        except Exception:
            logo = ""

        header = Table(
            [[header_copy, logo]],
            colWidths=[
                190 * mm,
                55 * mm,
            ],
        )

        header.setStyle(
            TableStyle(
                [
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "MIDDLE",
                    ),
                    (
                        "ALIGN",
                        (1, 0),
                        (1, 0),
                        "RIGHT",
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        0,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        0,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        0,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        0,
                    ),
                ]
            )
        )

        story.append(header)
        story.append(
            Spacer(
                1,
                4 * mm,
            )
        )

        # ====================================================
        # RESUMEN GENERAL - RESTAURADO
        # ====================================================

        summary_data = [
            [
                "Vencido",
                "Vence hoy",
                "Por vencer",
                "Total",
            ],
            [
                cls._money(
                    totals["overdue"]
                ),
                cls._money(
                    totals["due_today"]
                ),
                cls._money(
                    totals["upcoming"]
                ),
                cls._money(
                    totals["grand_total"]
                ),
            ],
        ]

        summary = Table(
            summary_data,
            colWidths=[
                58 * mm,
                58 * mm,
                58 * mm,
                58 * mm,
            ],
        )

        summary.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#F2F6F4"),
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#65726C"),
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "FONTNAME",
                        (0, 1),
                        (-1, 1),
                        "Helvetica-Bold",
                    ),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "RIGHT",
                    ),
                    (
                        "BOX",
                        (0, 0),
                        (-1, -1),
                        0.4,
                        colors.HexColor("#DDE6E1"),
                    ),
                    (
                        "INNERGRID",
                        (0, 0),
                        (-1, -1),
                        0.3,
                        colors.HexColor("#E8EEEB"),
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(summary)
        story.append(
            Spacer(
                1,
                5 * mm,
            )
        )

        # ====================================================
        # HELPERS
        # ====================================================

        def build_document_table(items):

            rows = [
                [
                    "Documento",
                    "Referencia",
                    "OC",
                    "Obra",
                    "Emisión",
                    "Vencimiento",
                    "Mora",
                    "Monto original",
                    "Saldo",
                ]
            ]

            for item in items:

                document_number = escape(
                    str(
                        item.get(
                            "document_number",
                            "",
                        )
                        or ""
                    )
                )

                document_cell = [
                    Paragraph(
                        document_number,
                        table_body_bold_style,
                    )
                ]

                reference_labels = []

                if item.get("is_refactored"):
                    refact_reference = str(
                        item.get(
                            "refacturation_reference"
                        )
                        or ""
                    ).strip()

                    if refact_reference:
                        reference_labels.append(
                            "Refacturaci&oacute;n de "
                            "<b>"
                            + escape(refact_reference)
                            + "</b>"
                        )
                    else:
                        reference_labels.append(
                            "Refacturaci&oacute;n"
                        )

                if item.get("is_claimed"):
                    reference_labels.append(
                        "Reclamada"
                    )

                if reference_labels:
                    situation_cell = Paragraph(
                        " &middot; ".join(
                            reference_labels
                        ),
                        table_body_style,
                    )
                else:
                    situation_cell = ""

                purchase_order = str(
                    item.get(
                        "purchase_order",
                        "",
                    )
                    or ""
                ).strip()

                work_reference = str(
                    item.get(
                        "work_reference",
                        "",
                    )
                    or ""
                ).strip()

                situation_cell = (
                    situation_cell
                    if situation_cell
                    else ""
                )

                purchase_order_cell = (
                    Paragraph(
                        escape(purchase_order),
                        table_body_style,
                    )
                    if purchase_order
                    else ""
                )

                work_reference_cell = (
                    Paragraph(
                        escape(work_reference),
                        table_body_style,
                    )
                    if work_reference
                    else ""
                )

                if (
                    item.get("category")
                    == "OVERDUE"
                ):
                    days = (
                        item.get(
                            "days_from_due"
                        )
                        or "-"
                    )
                else:
                    days = "-"

                rows.append(
                    [
                        document_cell,
                        situation_cell,
                        purchase_order_cell,
                        work_reference_cell,
                        Paragraph(
                            escape(
                                str(item["issue_date"])
                            ),
                            table_body_center_style,
                        ),
                        Paragraph(
                            escape(
                                str(item["due_date"])
                            ),
                            table_body_center_style,
                        ),
                        Paragraph(
                            escape(str(days)),
                            table_body_center_style,
                        ),
                        Paragraph(
                            escape(
                                cls._money(
                                    item[
                                        "original_amount"
                                    ]
                                )
                            ),
                            table_body_right_style,
                        ),
                        Paragraph(
                            escape(
                                cls._money(
                                    item[
                                        "balance_amount"
                                    ]
                                )
                            ),
                            table_body_right_bold_style,
                        ),
                    ]
                )

            table = Table(
                rows,
                repeatRows=1,
                colWidths=[
                    25 * mm,
                    48 * mm,
                    21 * mm,
                    44 * mm,
                    23 * mm,
                    24 * mm,
                    14 * mm,
                    23 * mm,
                    24 * mm,
                ],
            )

            table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#F1F5F3"),
                        ),
                        (
                            "TEXTCOLOR",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#5D6C64"),
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, 0),
                            "Helvetica-Bold",
                        ),
                        (
                            "FONTSIZE",
                            (0, 0),
                            (-1, 0),
                            7,
                        ),
                        (
                            "ALIGN",
                            (5, 1),
                            (-1, -1),
                            "RIGHT",
                        ),
                        (
                            "VALIGN",
                            (0, 0),
                            (-1, -1),
                            "TOP",
                        ),
                        (
                            "LINEBELOW",
                            (0, 0),
                            (-1, 0),
                            0.5,
                            colors.HexColor("#CDD9D3"),
                        ),
                        (
                            "LINEBELOW",
                            (0, 1),
                            (-1, -1),
                            0.2,
                            colors.HexColor("#E6ECE9"),
                        ),
                        (
                            "TOPPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                        (
                            "BOTTOMPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                    ]
                )
            )

            return table


        def build_subtotal(
            label,
            amount,
        ):

            subtotal_table = Table(
                [
                    [
                        label,
                        cls._money(amount),
                    ]
                ],
                colWidths=[
                    170 * mm,
                    45 * mm,
                ],
            )

            subtotal_table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, -1),
                            colors.HexColor("#F8FAF9"),
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, -1),
                            "Helvetica-Bold",
                        ),
                        (
                            "TEXTCOLOR",
                            (0, 0),
                            (-1, -1),
                            colors.HexColor("#34413B"),
                        ),
                        (
                            "ALIGN",
                            (1, 0),
                            (1, 0),
                            "RIGHT",
                        ),
                        (
                            "BOX",
                            (0, 0),
                            (-1, -1),
                            0.4,
                            colors.HexColor("#DDE6E1"),
                        ),
                        (
                            "TOPPADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                        (
                            "BOTTOMPADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                    ]
                )
            )

            return subtotal_table


        def add_section(
            *,
            title,
            items,
            subtotal,
            subtotal_label,
        ):

            if not items:
                return

            story.append(
                Paragraph(
                    title,
                    section_style,
                )
            )

            story.append(
                Spacer(
                    1,
                    2 * mm,
                )
            )

            story.append(
                build_document_table(
                    items
                )
            )

            story.append(
                build_subtotal(
                    subtotal_label,
                    subtotal,
                )
            )

            story.append(
                Spacer(
                    1,
                    5 * mm,
                )
            )

        # ====================================================
        # TRES GRUPOS CORRECTOS
        # ====================================================

        overdue_documents = [
            item
            for item in snapshot["documents"]
            if item.get("category")
            == "OVERDUE"
        ]

        due_today_documents = [
            item
            for item in snapshot["documents"]
            if item.get("category")
            == "DUE_TODAY"
        ]

        upcoming_documents = [
            item
            for item in snapshot["documents"]
            if item.get("category")
            == "UPCOMING"
        ]

        overdue_subtotal = sum(
            (
                Decimal(
                    str(
                        item.get(
                            "balance_amount",
                            0,
                        )
                    )
                )
                for item in overdue_documents
            ),
            Decimal("0"),
        )

        due_today_subtotal = sum(
            (
                Decimal(
                    str(
                        item.get(
                            "balance_amount",
                            0,
                        )
                    )
                )
                for item in due_today_documents
            ),
            Decimal("0"),
        )

        upcoming_subtotal = sum(
            (
                Decimal(
                    str(
                        item.get(
                            "balance_amount",
                            0,
                        )
                    )
                )
                for item in upcoming_documents
            ),
            Decimal("0"),
        )

        # ====================================================
        # VENCIDOS
        # ====================================================

        add_section(
            title="Documentos vencidos",
            items=overdue_documents,
            subtotal=overdue_subtotal,
            subtotal_label="Total documentos vencidos",
        )

        # ====================================================
        # VENCEN HOY
        # ====================================================

        add_section(
            title="Documentos que vencen hoy",
            items=due_today_documents,
            subtotal=due_today_subtotal,
            subtotal_label="Total documentos que vencen hoy",
        )

        # ====================================================
        # POR VENCER
        # ====================================================

        add_section(
            title="Documentos por vencer",
            items=upcoming_documents,
            subtotal=upcoming_subtotal,
            subtotal_label="Total documentos por vencer",
        )

        document.build(story)

        return output.getvalue()


    @classmethod
    def build_xlsx(cls, snapshot: dict) -> bytes:

        workbook = Workbook()

        summary_sheet = workbook.active
        summary_sheet.title = "Resumen"

        detail_sheet = workbook.create_sheet(
            "Detalle"
        )

        customer = snapshot["customer"]
        totals = snapshot["totals"]

        summary_rows = [
            ("Estado de Cuenta", ""),
            ("Cliente", customer["name"]),
            ("RUT", customer.get("rut") or ""),
            (
                "Fecha generación",
                snapshot["as_of_date"],
            ),
            (
                "Documentos",
                snapshot["document_count"],
            ),
            (
                "Total vencido",
                Decimal(totals["overdue"]),
            ),
            (
                "Vence hoy",
                Decimal(totals["due_today"]),
            ),
            (
                "Total por vencer",
                Decimal(totals["upcoming"]),
            ),
            (
                "Total general",
                Decimal(totals["grand_total"]),
            ),
        ]

        for row in summary_rows:
            summary_sheet.append(row)

        summary_sheet["A1"].font = Font(
            bold=True,
            size=16,
        )

        for row in range(2, 10):
            summary_sheet.cell(
                row=row,
                column=1,
            ).font = Font(bold=True)

        for row in range(6, 10):
            summary_sheet.cell(
                row=row,
                column=2,
            ).number_format = '#,##0'

        summary_sheet.column_dimensions[
            "A"
        ].width = 24

        summary_sheet.column_dimensions[
            "B"
        ].width = 45

        headers = [
            "Estado",
            "Tipo",
            "Documento",
            "Situación",
            "OC",
            "Obra",
            "Fecha emisión",
            "Fecha vencimiento",
            "Días en mora",
            "Condición de pago",
            "Monto original",
            "Saldo",
        ]

        detail_sheet.append(headers)

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="EAF3EF",
        )

        for cell in detail_sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(
                vertical="center"
            )

        for item in snapshot["documents"]:

            situation_parts = []

            if item.get("is_refactored"):
                refact_reference = str(
                    item.get(
                        "refacturation_reference"
                    )
                    or ""
                ).strip()

                if refact_reference:
                    situation_parts.append(
                        "Refacturación de "
                        + refact_reference
                    )
                else:
                    situation_parts.append(
                        "Refacturación"
                    )

            if item.get("is_claimed"):
                situation_parts.append(
                    "Reclamada"
                )

            situation_text = " · ".join(
                situation_parts
            )

            detail_sheet.append(
                [
                    item["category_label"],
                    item.get("document_type", ""),
                    item.get(
                        "document_number",
                        "",
                    ),
                    situation_text,
                    item.get(
                        "purchase_order",
                        "",
                    ),
                    item.get(
                        "work_reference",
                        "",
                    ),
                    item["issue_date"],
                    item["due_date"],
                    (
                        item.get(
                            "days_from_due",
                            0,
                        )
                        if item.get("category") == "OVERDUE"
                        else "-"
                    ),
                    item.get(
                        "payment_terms",
                        "",
                    ),
                    Decimal(
                        item["original_amount"]
                    ),
                    Decimal(
                        item["balance_amount"]
                    ),
                ]
            )

        detail_sheet.freeze_panes = "A2"
        detail_sheet.auto_filter.ref = (
            detail_sheet.dimensions
        )

        widths = [
            16,
            20,
            18,
            22,
            34,
            16,
            18,
            14,
            30,
            18,
            18,
        ]

        for index, width in enumerate(
            widths,
            start=1,
        ):
            detail_sheet.column_dimensions[
                get_column_letter(index)
            ].width = width

        for row in detail_sheet.iter_rows(
            min_row=2,
            min_col=10,
            max_col=11,
        ):
            for cell in row:
                cell.number_format = '#,##0'

        output = BytesIO()
        workbook.save(output)

        return output.getvalue()
